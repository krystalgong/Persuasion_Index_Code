"""Compliance-aware checks for optional Persuasion Index resources."""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import os
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import PI_score_generator as _scorer


OPTIONAL_RESOURCE_NAMES = (
    "spacy_model",
    "single_word_concreteness",
    "multiword_concreteness",
    "liwc",
    "nrc_vad",
)


class ResourceUnavailableError(RuntimeError):
    """Raised when the full feature set is requested without all resources."""


def _clean_env(name: str) -> str | None:
    value = os.environ.get(name, "").strip()
    return value or None


def _env_flag(name: str) -> bool:
    return os.environ.get(name, "").strip().lower() in {"1", "true", "yes", "on"}


def _resolve_path(env_name: str, default_name: str | Path) -> Path:
    configured = _clean_env(env_name)
    if configured:
        path = Path(configured).expanduser()
        if not path.is_absolute():
            path = _scorer.BASE_DIR / path
        return path.resolve()
    return (_scorer.get_helper_dir() / default_name).resolve()


def _package_version(name: str) -> str | None:
    try:
        return version(name)
    except PackageNotFoundError:
        return None


def _file_status(
    *,
    path: Path,
    configured_by: str,
    source_url: str,
    license_note: str,
    detail: str,
    features: tuple[str, ...],
) -> dict[str, Any]:
    available = path.is_file()
    digest = None
    read_error = None
    if available:
        try:
            hasher = hashlib.sha256()
            with path.open("rb") as handle:
                for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                    hasher.update(chunk)
            digest = hasher.hexdigest()
        except OSError as exc:
            available = False
            read_error = f"Could not read resource: {exc}"
    return {
        "available": available,
        "path": str(path),
        "configured_by": configured_by if _clean_env(configured_by) else None,
        "version": None,
        "sha256": digest,
        "detail": detail if available else (read_error or f"File not found: {path}"),
        "features": list(features),
        "source_url": source_url,
        "license_note": license_note,
    }


def _delimited_header(path: Path) -> set[str]:
    delimiter = "\t" if path.suffix.lower() in {".txt", ".tsv"} else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        return {str(value).strip() for value in next(reader, [])}


def _validate_columns(
    status: dict[str, Any],
    required: set[str],
) -> dict[str, Any]:
    if not status["available"]:
        return status

    path = Path(str(status["path"]))
    if path.suffix.lower() == ".xlsx":
        if importlib.util.find_spec("openpyxl") is None:
            status["available"] = False
            status["detail"] = (
                "Excel resource found, but openpyxl is not installed. "
                "Install persuasion-index[excel]."
            )
            return status
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            workbook = load_workbook(path, read_only=True, data_only=True)
            worksheet = workbook.active
            header = {
                str(value).strip()
                for value in next(worksheet.iter_rows(values_only=True), [])
                if value is not None
            }
            workbook.close()
        except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
            status["available"] = False
            status["detail"] = f"Could not read Excel resource header: {exc}"
            return status
        if not required <= header:
            status["available"] = False
            status["detail"] = (
                f"Expected columns {sorted(required)}; found {sorted(header)}"
            )
        return status

    try:
        header = _delimited_header(path)
    except (OSError, UnicodeError) as exc:
        status["available"] = False
        status["detail"] = f"Could not read resource header: {exc}"
        return status

    if not required <= header:
        status["available"] = False
        status["detail"] = (
            f"Expected columns {sorted(required)}; found {sorted(header)}"
        )
    return status


def _spacy_status() -> dict[str, Any]:
    disabled = _env_flag("PI_DISABLE_SPACY")
    spacy_installed = importlib.util.find_spec("spacy") is not None
    model_installed = importlib.util.find_spec("en_core_web_sm") is not None
    available = not disabled and spacy_installed and model_installed

    if disabled:
        detail = "Disabled by PI_DISABLE_SPACY."
    elif not spacy_installed:
        detail = "spaCy is not installed. Install persuasion-index[spacy]."
    elif not model_installed:
        detail = (
            "en_core_web_sm is not installed. Run: "
            "python -m spacy download en_core_web_sm"
        )
    else:
        detail = "spaCy and en_core_web_sm are installed."

    return {
        "available": available,
        "path": None,
        "configured_by": "PI_DISABLE_SPACY" if disabled else None,
        "version": {
            "spacy": _package_version("spacy"),
            "en_core_web_sm": _package_version("en-core-web-sm"),
        },
        "sha256": None,
        "detail": detail,
        "features": [
            "Evidence.named_entities",
            "Authority/Credibility.organizations",
        ],
        "source_url": "https://spacy.io/usage/models",
        "license_note": "Installed separately as versioned Python packages.",
    }


def check_resources() -> dict[str, dict[str, Any]]:
    """Return availability, provenance, and license notes for PI resources.

    This function never downloads a resource and never accepts third-party
    license terms on the user's behalf.
    """
    single = _file_status(
        path=_resolve_path(
            "PI_CONCRETENESS_FILE",
            "Brysbaert_concretness_dataset.xlsx",
        ),
        configured_by="PI_CONCRETENESS_FILE",
        source_url="https://biblio.ugent.be/publication/5774089",
        license_note="Obtain from the official source; not redistributed by PI.",
        detail="Single-word concreteness ratings found.",
        features=("Specificity.lexical_concreteness",),
    )
    single = _validate_columns(single, {"Word", "Conc.M"})

    multiword = _file_status(
        path=_resolve_path(
            "PI_MWE_CONCRETENESS_FILE",
            "MultiwordExpression_Concreteness_Ratings.csv",
        ),
        configured_by="PI_MWE_CONCRETENESS_FILE",
        source_url="https://osf.io/ksypa/",
        license_note="Obtain from the authors' source; not redistributed by PI.",
        detail="Multiword-expression concreteness ratings found.",
        features=("Specificity.lexical_concreteness",),
    )
    multiword = _validate_columns(multiword, {"Expression", "Mean_C"})

    liwc = _file_status(
        path=_resolve_path("PI_LIWC_FILE", "en_liwc.txt"),
        configured_by="PI_LIWC_FILE",
        source_url="https://www.liwc.app/download",
        license_note=(
            "A valid LIWC license is required. Do not redistribute the "
            "dictionary. PI supports legacy .dic, not LIWC-22 .dicx."
        ),
        detail="Licensed LIWC-compatible dictionary found.",
        features=(
            "Specificity.interactional_immediacy",
            "Sentiment.anger",
            "Sentiment.sadness",
            "Sentiment.fear_threat",
            "Sentiment.joy_gain",
            "Engagement.identification",
            "Engagement.self_reference",
            "Engagement.past",
            "Engagement.imagery",
            "Engagement.characters",
        ),
    )
    if liwc["available"] and Path(str(liwc["path"])).suffix.lower() == ".dicx":
        liwc["available"] = False
        liwc["detail"] = (
            "LIWC-22 .dicx is not supported by the current loader. Do not "
            "convert or extract licensed data; use a supported licensed .dic "
            "or a future official LIWC CLI integration."
        )

    nrc = _file_status(
        path=_resolve_path(
            "PI_NRC_VAD_FILE",
            Path("NRC-VAD-Lexicon-v2.1")
            / "Unigrams"
            / "unigrams-NRC-VAD-Lexicon-v2.1.txt",
        ),
        configured_by="PI_NRC_VAD_FILE",
        source_url="https://saifmohammad.com/WebPages/nrc-vad.html",
        license_note=(
            "Non-commercial research/educational use; citation required; "
            "redistribution prohibited by the provider."
        ),
        detail="NRC-VAD v2.1 unigram resource found.",
        features=(
            "Sentiment.valence",
            "Sentiment.arousal",
            "Sentiment.dominance",
        ),
    )
    nrc = _validate_columns(nrc, {"term", "valence", "arousal", "dominance"})

    return {
        "spacy_model": _spacy_status(),
        "single_word_concreteness": single,
        "multiword_concreteness": multiword,
        "liwc": liwc,
        "nrc_vad": nrc,
    }


def missing_resources() -> list[str]:
    """Return optional feature resources that are currently unavailable."""
    statuses = check_resources()
    return [name for name in OPTIONAL_RESOURCE_NAMES if not statuses[name]["available"]]


def require_resources() -> dict[str, dict[str, Any]]:
    """Require every optional resource used by the full feature set."""
    statuses = check_resources()
    missing = [
        name for name in OPTIONAL_RESOURCE_NAMES if not statuses[name]["available"]
    ]
    if missing:
        joined = ", ".join(missing)
        raise ResourceUnavailableError(
            "The full feature set requires additional local resources: "
            f"{joined}. Run `persuasion-index doctor` for setup details."
        )
    return statuses
